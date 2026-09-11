from pathlib import Path

import docx
import pymupdf

from app.ingestion.extract import extract_docx_text, extract_from_file, extract_pdf_text, extract_plain_text
from app.schemas.document import SourceFormat


def _make_pdf(tmp_path: Path, text: str) -> Path:
    path = tmp_path / "doc.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), text)
    pdf.save(str(path))
    pdf.close()
    return path


def _make_docx(tmp_path: Path, paragraphs: list[str]) -> Path:
    path = tmp_path / "doc.docx"
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    document.save(str(path))
    return path


def test_extract_pdf_text(tmp_path):
    pdf_path = _make_pdf(tmp_path, "Jane Doe - Software Engineer")
    result = extract_pdf_text(pdf_path)
    assert result.source_format == SourceFormat.PDF
    assert "Jane Doe" in result.raw_text
    assert result.page_count == 1


def test_extract_pdf_warns_on_near_empty_text(tmp_path):
    path = tmp_path / "blank.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(str(path))
    pdf.close()
    result = extract_pdf_text(path)
    assert result.extraction_warnings


def test_extract_docx_text_includes_tables(tmp_path):
    docx_path = _make_docx(tmp_path, ["Experience", "Software Engineer at Acme"])
    document = docx.Document(str(docx_path))
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Python"
    document.save(str(docx_path))

    result = extract_docx_text(docx_path)
    assert "Software Engineer at Acme" in result.raw_text
    assert "Python" in result.raw_text


def test_extract_plain_text_strips_and_flags_empty():
    result = extract_plain_text("  hello world  ")
    assert result.raw_text == "hello world"
    assert result.extraction_warnings == []

    empty = extract_plain_text("   ")
    assert empty.extraction_warnings


def test_extract_from_file_rejects_unsupported_extension(tmp_path):
    bad_file = tmp_path / "resume.rtf"
    bad_file.write_text("hello")
    try:
        extract_from_file(bad_file)
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_extract_from_file_dispatches_by_extension(tmp_path):
    pdf_path = _make_pdf(tmp_path, "hello")
    result = extract_from_file(pdf_path)
    assert result.source_format == SourceFormat.PDF


def test_extract_from_file_reads_txt(tmp_path):
    txt_path = tmp_path / "jd.txt"
    txt_path.write_text("Backend Engineer role requiring Python.")
    result = extract_from_file(txt_path)
    assert result.source_format == SourceFormat.TEXT
    assert "Python" in result.raw_text
